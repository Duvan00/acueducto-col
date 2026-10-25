"""Pruebas de integración contra PostgreSQL/PostGIS (base acueducto_test)."""
import io
import json
import shutil
import subprocess

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import semillas
from app.db import conexion
from app.main import app


@pytest.fixture(scope="module")
def cli():
    semillas.inicializar(demo=True, reiniciar=True)
    return TestClient(app)


def uid(codigo):
    with conexion() as con:
        return con.execute("SELECT id FROM unidad_territorial WHERE codigo=%s", (codigo,)).fetchone()["id"]


def test_territorio(cli):
    d = cli.get("/api/departamentos").json()
    assert len([x for x in d if not x["es_demostracion"]]) == 33
    m = cli.get("/api/municipios", params={"dpto": "05"}).json()
    assert any(x["codigo"] == "05001" for x in m)
    assert cli.get("/api/buscar", params={"q": "medellin"}).json()[0]["codigo"] == "05001"   # sin tilde
    assert cli.get("/api/buscar", params={"q": "05001"}).json()[0]["nombre"] == "MEDELLÍN"


def test_importar_proyecciones_formato_largo_y_versionado(cli):
    def archivo(valor):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Proyecciones de población municipal por área"])   # título previo al encabezado
        ws.append([])
        ws.append(["DP", "DPNOM", "DPMP", "MPIO", "AÑO", "ÁREA GEOGRÁFICA", "Población"])
        for a in range(2018, 2043):
            ws.append(["05", "Antioquia", "05001", "Medellín", a, "Total", valor + a])
            ws.append(["05", "Antioquia", "05001", "Medellín", a, "Cabecera Municipal", valor])
            ws.append(["05", "Antioquia", "05001", "Medellín", a, "Centros Poblados y Rural Disperso", a])
        b = io.BytesIO(); wb.save(b); return b.getvalue()
    r = cli.post("/api/importar/proyecciones", data={"version": "TEST-v1"},
                 files={"archivo": ("proy.xlsx", archivo(1000), "application/octet-stream")}).json()
    assert r["registros"] == 75
    r2 = cli.post("/api/importar/proyecciones", data={"version": "TEST-v2"},
                  files={"archivo": ("proy.xlsx", archivo(2000), "application/octet-stream")}).json()
    assert r2["registros"] == 75
    with conexion() as con:
        filas = con.execute("""SELECT serie, vigente, count(*) n FROM dato_demografico d JOIN unidad_territorial u ON u.id=d.unidad_id
                               WHERE u.codigo='05001' GROUP BY 1,2 ORDER BY 1""").fetchall()
    assert {(f["serie"], f["vigente"]) for f in filas} == {("PROY_MUN_AREA_TEST-v1", False), ("PROY_MUN_AREA_TEST-v2", True)}
    v = cli.get("/api/validaciones", params={"municipio": "05001"}).json()
    assert next(x for x in v if x["clave"] == "suma_areas")["n"] == 0


def test_importar_censo_formato_csv(cli):
    csv = "Código;Área;Personas;Hombres;Mujeres;Viviendas;Hogares\n05001;Total;1.000.000;480.000;520.000;350.000;340.000\n05001;Cabecera;900.000;;;;\n"
    r = cli.post("/api/importar/censo", data={"anio": 2018, "serie": "CNPV2018_TEST", "producto": "CNPV 2018 prueba"},
                 files={"archivo": ("c.csv", csv.encode(), "text/csv")}).json()
    assert r["registros"] == 6
    d = cli.get(f"/api/unidades/{uid('05001')}/demografia").json()
    c = next(x for x in d["censos"] if x["area"] == "total")
    assert c["poblacion"] == 1_000_000 and c["clasificacion"] == "censo_oficial_dane"


def test_no_se_importan_estimaciones_ni_censos_de_otra_entidad(cli):
    r = cli.post("/api/importar/censo", data={"anio": 2018, "serie": "X", "producto": "x", "entidad": "Gobernación"},
                 files={"archivo": ("c.csv", b"codigo,poblacion\n05001,1\n", "text/csv")})
    assert r.status_code == 422
    r = cli.post("/api/importar/censo", data={"anio": 2018, "serie": "X", "producto": "x", "clasificacion": "estimacion_tecnica"},
                 files={"archivo": ("c.csv", b"codigo,poblacion\n05001,1\n", "text/csv")})
    assert r.status_code == 422


def test_estimacion_no_puede_ser_oficial():
    import psycopg
    with pytest.raises(psycopg.errors.CheckViolation):
        with conexion() as con:
            con.execute("""INSERT INTO dato_demografico (unidad_id, anio, variable, valor, clasificacion, fuente_id, serie)
                           VALUES (%s, 2020, 'poblacion', 5, 'estimacion_tecnica', 1, 'CNPV2018')""", (uid("05001"),))


def test_dato_campo_requiere_soporte(cli):
    r = cli.post(f"/api/unidades/{uid('V-XX001-02')}/datos", json={"anio": 2025, "variable": "poblacion", "valor": 300, "soporte": " "})
    assert r.status_code == 422


REQ = lambda: {
    "proyecto": {"nombre": "Prueba", "tipo": "regional"}, "anio_inicial": 2026, "periodo_diseno": 25,
    "localidades": [{"unidad_id": uid("XX001"), "area": "cabecera", "metodo": "geometrico"},
                    {"unidad_id": uid("XX001001"), "metodo": "aritmetico", "cobertura_pct": 90, "anio_ingreso": 2030},
                    {"unidad_id": uid("V-XX001-01"), "metodo": "geometrico", "tasa_referencia": "municipio_area"}],
    "demandas": [{"tipo": "institucional", "caudal_lps": 0.3, "justificacion": "Colegio"}],
    "etapas": [{"nombre": "E1", "anio_inicio": 2026, "anio_fin": 2038}],
}


def test_calculo_integral(cli):
    r = cli.post("/api/calcular", json=REQ()).json()
    assert r["es_demostracion"]
    assert r["advertencias"][0]["codigo"] == "DATOS_DEMOSTRACION"
    loc = r["localidades"]
    assert loc[2]["seleccion"]["poblacion_base"] == 412            # dato de campo como base
    assert loc[2]["tasa_referencia"] == pytest.approx(((3550 / 3300) ** (1 / 12) - 1 + (4080 / 3550) ** (1 / 13) - 1) / 2)
    s = {p["anio"]: p["poblacion_servida"] for p in r["sistema"]["serie_total"]}
    assert s[2051] == pytest.approx(r["caudales"]["poblacion"])
    q = r["caudales"]
    assert q["dotacion_neta"] == 130 and q["k1"] == 1.2           # >12.500 hab, 1790 m s.n.m.


def test_doble_conteo_municipio_y_vereda(cli):
    req = REQ()
    req["localidades"][0]["area"] = "total"
    r = cli.post("/api/calcular", json=req).json()
    assert any(a["codigo"] == "JERARQUIA_SUPERPUESTA" for a in r["advertencias"])
    req["localidades"][0]["area"] = "cabecera"
    r = cli.post("/api/calcular", json=req).json()
    assert not any(a["codigo"] == "JERARQUIA_SUPERPUESTA" for a in r["advertencias"])


def test_area_de_influencia(cli):
    g = {"type": "Polygon", "coordinates": [[[-79.57, 4.05], [-79.49, 4.05], [-79.49, 4.14], [-79.57, 4.14], [-79.57, 4.05]]]}
    r = cli.post("/api/geo/localidades_en_area", json={"geometria": g}).json()
    nombres = {x["nombre"] for x in r}
    assert "Vereda El Roble (FICTICIA)" in nombres and "Centro Poblado El Cruce (FICTICIO)" in nombres


def test_lectura_kml(cli):
    kml = b"""<?xml version="1.0"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document><Placemark><name>A</name>
      <Polygon><outerBoundaryIs><LinearRing><coordinates>-79.6,4 -79.5,4 -79.5,4.1 -79.6,4</coordinates></LinearRing></outerBoundaryIs></Polygon>
      </Placemark></Document></kml>"""
    r = cli.post("/api/geo/leer", files={"archivo": ("a.kml", kml, "application/xml")}).json()
    assert r["features"][0]["geometry"]["type"] == "MultiPolygon"


@pytest.mark.parametrize("fmt", ["docx", "pdf", "xlsx"])
def test_memorias(cli, fmt):
    r = cli.post(f"/api/memoria/{fmt}", json=REQ())
    assert r.status_code == 200 and len(r.content) > 5000


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice no disponible")
def test_excel_formulas_reproducen_calculo(cli, tmp_path):
    res = cli.post("/api/calcular", json=REQ()).json()
    x = cli.post("/api/memoria/xlsx", json=REQ()).content
    (tmp_path / "m.xlsx").write_bytes(x)
    out = tmp_path / "out"
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(tmp_path / "m.xlsx")],
                   check=True, capture_output=True, timeout=180)
    ws = openpyxl.load_workbook(out / "m.xlsx", data_only=True)["Caudales"]
    v = {ws.cell(i, 2).value: ws.cell(i, 3).value for i in range(4, 40) if isinstance(ws.cell(i, 2).value, str)}
    q = res["caudales"]
    assert v["P"] == pytest.approx(q["poblacion"], rel=1e-9)
    assert v["Qmd"] == pytest.approx(q["qmd_lps"], rel=1e-9)
    assert v["QMD"] == pytest.approx(q["QMD_lps"], rel=1e-9)
    assert v["QMH"] == pytest.approx(q["QMH_lps"], rel=1e-9)
