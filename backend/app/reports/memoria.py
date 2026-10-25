"""Generación de la memoria de cálculo en Word, PDF y Excel (con fórmulas vivas)."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..calc.proyeccion import NOMBRE_METODO
from ..servicio import ETIQUETA_CLASIF
from . import contenido

MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def generar(formato: str, res: dict) -> tuple[bytes, str]:
    return {"docx": docx, "pdf": pdf, "xlsx": xlsx}[formato](res), MIME[formato]


# ---------------------------------------------------------------- Word

def docx(res: dict) -> bytes:
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.2)
    t = d.add_heading("Memoria de cálculo — Población de diseño y caudales de acueducto", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph(res["proyecto"]["nombre"]).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for b in contenido.construir(res):
        k = b[0]
        if k == "h1":
            d.add_heading(b[1], 1)
        elif k == "h2":
            d.add_heading(b[1], 2)
        elif k == "p":
            d.add_paragraph(b[1])
        elif k == "formula":
            par = d.add_paragraph()
            r = par.add_run(b[1])
            r.font.name, r.font.size, r.italic = "Cambria Math", Pt(10), True
            par.paragraph_format.left_indent = Cm(1)
        elif k in ("aviso", "aviso_suave"):
            par = d.add_paragraph()
            r = par.add_run(b[1])
            r.bold = k == "aviso"
            r.font.color.rgb = RGBColor(0xB0, 0x00, 0x20) if k == "aviso" else RGBColor(0x8A, 0x5A, 0x00)
        elif k == "tabla":
            enc, filas = b[1], b[2]
            tb = d.add_table(rows=1, cols=len(enc))
            tb.style = "Light Grid Accent 1"
            for i, h in enumerate(enc):
                tb.rows[0].cells[i].text = str(h)
            for f in filas:
                cells = tb.add_row().cells
                for i, v in enumerate(f):
                    cells[i].text = str(v)
            for row in tb.rows:
                for c in row.cells:
                    for par in c.paragraphs:
                        for r in par.runs:
                            r.font.size = Pt(8)
            d.add_paragraph()
        elif k == "imagen":
            d.add_picture(io.BytesIO(b[1]), width=Cm(b[2]))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

def pdf(res: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                            bottomMargin=2 * cm, title="Memoria de cálculo", author=res["proyecto"].get("elaboro") or "")
    ss = getSampleStyleSheet()
    sty = {
        "t": ParagraphStyle("t", parent=ss["Title"], fontSize=15),
        "h1": ParagraphStyle("h1", parent=ss["Heading2"], textColor=colors.HexColor("#0b3c5d")),
        "h2": ParagraphStyle("h2", parent=ss["Heading3"]),
        "p": ParagraphStyle("p", parent=ss["BodyText"], fontSize=9, leading=12),
        "f": ParagraphStyle("f", parent=ss["BodyText"], fontName="Times-Italic", fontSize=9.5, leftIndent=18),
        "a": ParagraphStyle("a", parent=ss["BodyText"], fontSize=10, textColor=colors.HexColor("#b00020"),
                            fontName="Helvetica-Bold"),
        "as": ParagraphStyle("as", parent=ss["BodyText"], fontSize=8, textColor=colors.HexColor("#8a5a00")),
        "c": ParagraphStyle("c", parent=ss["BodyText"], fontSize=7, leading=8.5),
    }
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    E = [Paragraph("Memoria de cálculo — Población de diseño y caudales de acueducto", sty["t"]),
         Paragraph(esc(res["proyecto"]["nombre"]), ss["Heading3"])]
    ancho = letter[0] - 4 * cm
    for b in contenido.construir(res):
        k = b[0]
        if k in ("h1", "h2", "p"):
            E.append(Paragraph(esc(b[1]), sty[k]))
        elif k == "formula":
            E.append(Paragraph(esc(b[1]), sty["f"]))
        elif k == "aviso":
            E.append(Paragraph(esc(b[1]), sty["a"]))
        elif k == "aviso_suave":
            E.append(Paragraph(esc(b[1]), sty["as"]))
        elif k == "tabla":
            enc, filas, anchos = b[1], b[2], b[3]
            data = [[Paragraph(f"<b>{esc(h)}</b>", sty["c"]) for h in enc]] + \
                   [[Paragraph(esc(v), sty["c"]) for v in f] for f in filas]
            cw = [ancho * a / sum(anchos) for a in anchos] if anchos else [ancho / len(enc)] * len(enc)
            t = Table(data, colWidths=cw, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .3, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbe9f4")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            E += [t, Spacer(1, 6)]
        elif k == "imagen":
            img = Image(io.BytesIO(b[1]), width=b[2] * cm, height=b[2] * cm * 0.53)
            E.append(KeepTogether([img, Spacer(1, 4)]))

    def pie(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        txt = f"{res['proyecto']['nombre']} — motor v{res['version_motor']} — {res['fecha']} — pág. {doc_.page}"
        if res["es_demostracion"]:
            txt = "DATOS FICTICIOS DE DEMOSTRACIÓN — " + txt
        canvas.drawString(2 * cm, 1.2 * cm, txt)
        canvas.restoreState()

    doc.build(E, onFirstPage=pie, onLaterPages=pie)
    return buf.getvalue()


# ---------------------------------------------------------------- Excel (fórmulas vivas)

AZUL = PatternFill("solid", fgColor="DBE9F4")
AMAR = PatternFill("solid", fgColor="FFF4CC")     # celdas de entrada editables
BOLD = Font(bold=True)
FINO = Border(*(Side(style="thin", color="BBBBBB"),) * 4)


def _enc(ws, fila, textos, col=1):
    for i, t in enumerate(textos):
        c = ws.cell(fila, col + i, t)
        c.font, c.fill, c.border = BOLD, AZUL, FINO
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _nombre(wb, nombre, ws, ref):
    wb.defined_names[nombre] = DefinedName(nombre, attr_text=f"'{ws.title}'!{ref}")


def xlsx(res: dict) -> bytes:
    wb = Workbook()
    q = res["caudales"]
    a0, af = res["anio_inicial"], res["anio_final"]
    anios = list(range(a0, af + 1))

    # --- Hoja Proyecto
    ws = wb.active
    ws.title = "Proyecto"
    ws["A1"] = "Memoria de cálculo — Población de diseño y caudales de acueducto"
    ws["A1"].font = Font(bold=True, size=14)
    if res["es_demostracion"]:
        ws["A2"] = "DATOS FICTICIOS DE DEMOSTRACIÓN — NO VÁLIDO PARA DISEÑO"
        ws["A2"].font = Font(bold=True, color="B00020")
    pr = res["proyecto"]
    filas = [("Proyecto", pr["nombre"]), ("Entidad", pr["entidad"]), ("Elaboró", pr["elaboro"]), ("Tipo", pr["tipo"]),
             ("Fecha", res["fecha"]), ("Motor", res["version_motor"]), ("SHA-256 entrada", res["sha256"]),
             ("Año inicial", a0), ("Período de diseño (años)", res["periodo_diseno"]), ("Año horizonte", None)]
    for i, (k, v) in enumerate(filas, start=4):
        ws.cell(i, 1, k).font = BOLD
        ws.cell(i, 2, v)
    ws["B11"].fill = ws["B12"].fill = AMAR
    ws["B13"] = "=B11+B12"
    _nombre(wb, "AnioInicial", ws, "$B$11")
    _nombre(wb, "AnioFinal", ws, "$B$13")
    r = 15
    ws.cell(r, 1, "Celdas amarillas: datos de entrada editables. Las demás contienen fórmulas.").font = Font(italic=True)
    r += 2
    ws.cell(r, 1, "Fuentes").font = BOLD
    for s in res["fuentes"]:
        r += 1
        ws.cell(r, 1, s)
    r += 2
    ws.cell(r, 1, "Advertencias").font = BOLD
    for a in res["advertencias"]:
        r += 1
        ws.cell(r, 1, a["nivel"].upper())
        ws.cell(r, 2, a["mensaje"])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110

    # --- Hojas por localidad
    hojas_loc = []
    for idx, loc in enumerate(res["localidades"], start=1):
        wl = wb.create_sheet(f"Loc{idx}")
        hojas_loc.append(wl.title)
        s = loc["seleccion"]
        wl["A1"] = loc["nombre"]
        wl["A1"].font = Font(bold=True, size=12)
        wl["A2"] = f"Área: {contenido.AREA[loc['area']]} — Método adoptado: {NOMBRE_METODO[loc['metodo']]}"
        _enc(wl, 4, ["Año censo", "Población", "Incluir (1/0)", "Fuente", "Clasificación"])
        censos = sorted(loc["censos"], key=lambda c: c["anio"])
        for i, c in enumerate(censos, start=5):
            wl.cell(i, 1, c["anio"]).fill = AMAR
            wl.cell(i, 2, c["poblacion"]).fill = AMAR
            wl.cell(i, 3, 1 if c["comparable"] else 0).fill = AMAR
            wl.cell(i, 4, c["fuente"])
            wl.cell(i, 5, ETIQUETA_CLASIF.get(c["clasificacion"], c["clasificacion"]))
        fin_c = 4 + len(censos)
        # Tasas por período (fila i usa censos i e i+1)
        rt = fin_c + 2
        _enc(wl, rt, ["Período", "ka (hab/año)", "r geométrica", "k exponencial", "Peso (comparable)"])
        filas_t = []
        for j in range(len(censos) - 1):
            i1, i2, ri = 5 + j, 6 + j, rt + 1 + j
            wl.cell(ri, 1, f'=A{i1}&"–"&A{i2}')
            wl.cell(ri, 2, f"=(B{i2}-B{i1})/(A{i2}-A{i1})")
            wl.cell(ri, 3, f"=(B{i2}/B{i1})^(1/(A{i2}-A{i1}))-1")
            wl.cell(ri, 4, f"=(LN(B{i2})-LN(B{i1}))/(A{i2}-A{i1})")
            wl.cell(ri, 5, f"=C{i1}*C{i2}")
            filas_t.append(ri)
        ra = rt + len(censos) + 1
        wl.cell(ra, 1, "Tasa adoptada").font = BOLD
        if filas_t:
            t1, t2 = filas_t[0], filas_t[-1]
            for col in "BCD":
                if res["criterio_tasa"] == "ultimo_periodo":
                    # último período con peso 1
                    wl[f"{col}{ra}"] = (f"=LOOKUP(2,1/(E{t1}:E{t2}=1),{col}{t1}:{col}{t2})")
                else:
                    wl[f"{col}{ra}"] = f"=SUMPRODUCT({col}{t1}:{col}{t2},E{t1}:E{t2})/SUM(E{t1}:E{t2})"
        # Parámetros de la localidad
        rp = ra + 2
        es_manual = s["metodo"] != "dane_oficial" and loc.get("comparacion", {}).get(s["metodo"], {}).get("tasa") != s["tasa"]
        params = [("Año base (Tb)", s["anio_base"], True), ("Población base (Pb)", s["poblacion_base"], True),
                  ("Cobertura (%)", loc["cobertura_pct"], True), ("Año de ingreso al sistema", loc["anio_ingreso"] or a0, True),
                  ("Método (1 aritm., 2 geom., 3 exp., 4 DANE)",
                   {"aritmetico": 1, "geometrico": 2, "exponencial": 3, "dane_oficial": 4}[s["metodo"]], True),
                  ("Tasa manual del método adoptado (vacío = intercensal)", s["tasa"] if es_manual else None, True)]
        for i, (k, v, ent) in enumerate(params):
            wl.cell(rp + i, 1, k).font = BOLD
            c = wl.cell(rp + i, 2, v)
            if ent:
                c.fill = AMAR
        Tb, Pb, Cob, Ing, Met, Tm = (f"$B${rp + i}" for i in range(6))
        met_sel = {"aritmetico": 1, "geometrico": 2, "exponencial": 3}.get(s["metodo"])

        def tasa_col(m, col_tasa):
            # La tasa manual sólo se aplica al método adoptado; los demás usan la tasa intercensal.
            if met_sel == m and es_manual:
                return f'IF({Tm}<>"",{Tm},${col_tasa}${ra})' if filas_t else Tm
            return f"${col_tasa}${ra}" if filas_t else None
        ka, rg, ke = tasa_col(1, "B"), tasa_col(2, "C"), tasa_col(3, "D")
        # Tabla de proyección
        rpj = rp + 8
        _enc(wl, rpj, ["Año", "Aritmético", "Geométrico", "Exponencial", "DANE oficial / extrap.", "Clasificación DANE",
                       "Adoptada", "Servida"])
        serie_d = {p_["anio"]: p_ for p_ in loc["comparacion"].get("dane_oficial", {}).get("serie", [])}
        for i, t in enumerate(anios):
            rr = rpj + 1 + i
            wl.cell(rr, 1, "=AnioInicial" if i == 0 else f"=A{rr - 1}+1")
            if ka:
                wl.cell(rr, 2, f"={Pb}+{ka}*(A{rr}-{Tb})")
            if rg:
                wl.cell(rr, 3, f"={Pb}*(1+{rg})^(A{rr}-{Tb})")
            if ke:
                wl.cell(rr, 4, f"={Pb}*EXP({ke}*(A{rr}-{Tb}))")
            pd = serie_d.get(t)
            if pd and pd["poblacion"] is not None:
                wl.cell(rr, 5, pd["poblacion"])
                wl.cell(rr, 6, ETIQUETA_CLASIF.get(pd["clasificacion"], pd["clasificacion"]))
            wl.cell(rr, 7, f"=CHOOSE({Met},B{rr},C{rr},D{rr},E{rr})")
            wl.cell(rr, 8, f"=IF(A{rr}>={Ing},G{rr}*{Cob}/100,0)")
            for col in range(2, 9):
                if col != 6:
                    wl.cell(rr, col).number_format = "#,##0"
        loc["_rango_servida"] = (rpj + 1, rpj + len(anios))
        for col, w in zip("ABCDEFGH", (34, 14, 14, 14, 18, 30, 14, 14)):
            wl.column_dimensions[col].width = w

    # --- Sistema
    wsy = wb.create_sheet("Sistema")
    _enc(wsy, 1, ["Año"] + [l["nombre"][:30] for l in res["localidades"]] + ["Total servido"])
    for i, t in enumerate(anios):
        rr = 2 + i
        wsy.cell(rr, 1, f"='{hojas_loc[0]}'!A{res['localidades'][0]['_rango_servida'][0] + i}")
        refs = []
        for j, (h, loc) in enumerate(zip(hojas_loc, res["localidades"])):
            fila = loc["_rango_servida"][0] + i
            wsy.cell(rr, 2 + j, f"='{h}'!H{fila}").number_format = "#,##0"
            refs.append(f"{get_column_letter(2 + j)}{rr}")
        tc = 2 + len(hojas_loc)
        wsy.cell(rr, tc, f"=SUM({refs[0]}:{refs[-1]})").number_format = "#,##0"
    fila_final = 1 + len(anios)
    col_tot = get_column_letter(2 + len(hojas_loc))
    _nombre(wb, "PoblacionDiseno", wsy, f"${col_tot}${fila_final}")
    for l in res["localidades"]:
        l.pop("_rango_servida", None)

    # --- Caudales
    wq = wb.create_sheet("Caudales")
    wq["A1"] = "Cálculo de caudales — Res. 0330/2017 arts. 43, 44 y 47"
    wq["A1"].font = Font(bold=True, size=12)
    _enc(wq, 3, ["Parámetro", "Símbolo", "Valor", "Unidad", "Fórmula / fundamento"])
    filas_q = [
        ("Población de diseño", "P", "=PoblacionDiseno", "hab", "Hoja Sistema, año horizonte"),
        ("Altitud promedio zona atendida", "h", q.get("altitud_msnm"), "m s.n.m.", "Entrada"),
        ("Dotación neta", "dneta", q["dotacion_neta"], "L/hab·día", f"Máx. normativa: {q['dotacion_neta_max']} (art. 43)"),
        ("Pérdidas técnicas", "%p", q["perdidas_pct"] / 100, "-", "Máx. 25 % (art. 44)"),
        ("Dotación bruta", "dbruta", "=C6/(1-C7)", "L/hab·día", "dbruta = dneta/(1-%p)"),
        ("Caudal medio diario residencial", "Qmd,res", "=C4*C8/86400", "L/s", "P·dbruta/86400"),
        ("K1", "K1", q["k1"], "-", "Art. 47 (máx. 1,3 si P≤12.500; 1,2 si P>12.500)"),
        ("K2", "K2", q["k2"], "-", "Art. 47 (máx. 1,6 si P≤12.500; 1,5 si P>12.500)"),
    ]
    for i, (a, b_, c, d_, e) in enumerate(filas_q, start=4):
        wq.cell(i, 1, a)
        wq.cell(i, 2, b_)
        cc = wq.cell(i, 3, c)
        wq.cell(i, 4, d_)
        wq.cell(i, 5, e)
        if not (isinstance(c, str) and c.startswith("=")):
            cc.fill = AMAR
    wq["C7"].number_format = "0.0%"
    # Demandas adicionales
    rd = 14
    _enc(wq, rd, ["Demanda adicional", "Caudal declarado (L/s)", "Incluye pérdidas (1/0)", "Afecta K1 (1/0)", "Afecta K2 (1/0)",
                  "Caudal bruto (L/s)", "Aporte QMD (L/s)", "Aporte QMH (L/s)", "Justificación"])
    dem = q["demandas"] or []
    n_d = max(len(dem), 1)
    for i in range(n_d):
        rr = rd + 1 + i
        d_ = dem[i] if i < len(dem) else {"tipo": "(ninguna)", "caudal_lps": 0, "incluye_perdidas": False,
                                         "afecta_k1": True, "afecta_k2": True, "justificacion": ""}
        wq.cell(rr, 1, d_["tipo"])
        for col, v in ((2, d_["caudal_lps"]), (3, int(d_["incluye_perdidas"])), (4, int(d_["afecta_k1"])),
                       (5, int(d_["afecta_k2"]))):
            wq.cell(rr, col, v).fill = AMAR
        wq.cell(rr, 6, f"=IF(C{rr}=1,B{rr},B{rr}/(1-$C$7))")
        wq.cell(rr, 7, f"=F{rr}*IF(D{rr}=1,$C$10,1)")
        wq.cell(rr, 8, f"=G{rr}*IF(E{rr}=1,$C$11,1)")
        wq.cell(rr, 9, d_["justificacion"])
    d1, d2 = rd + 1, rd + n_d
    rr = d2 + 2
    _enc(wq, rr, ["Resultado", "Símbolo", "L/s", "m³/día", "Fórmula"])
    res_q = [("Caudal medio diario", "Qmd", f"=C9+SUM(F{d1}:F{d2})", "Qmd = Qmd,res + ΣQad"),
             ("Caudal máximo diario", "QMD", f"=C10*C9+SUM(G{d1}:G{d2})", "QMD = K1·Qmd"),
             ("Caudal máximo horario", "QMH", f"=C11*C10*C9+SUM(H{d1}:H{d2})", "QMH = K2·QMD")]
    for i, (a, b_, c, e) in enumerate(res_q, start=rr + 1):
        wq.cell(i, 1, a).font = BOLD
        wq.cell(i, 2, b_)
        wq.cell(i, 3, c).number_format = "0.000"
        wq.cell(i, 4, f"=C{i}*86.4").number_format = "#,##0.0"
        wq.cell(i, 5, e)
    if q["demandas_excluidas"]:
        wq.cell(rr + 5, 1, "Excluidas por estar contenidas en la dotación: " + ", ".join(q["demandas_excluidas"]))
    for col, w in zip("ABCDEFGHI", (34, 20, 16, 14, 40, 16, 16, 16, 50)):
        wq.column_dimensions[col].width = w

    # --- Normativa
    wn = wb.create_sheet("Normativa")
    _enc(wn, 1, ["Clave", "Valor", "Unidad", "Condición", "Norma", "Artículo", "Texto de referencia", "Verificación"])
    for i, x in enumerate(res["parametros_normativos"], start=2):
        for j, k in enumerate(["clave", "valor", "unidad", "condicion", "norma", "articulo", "texto_referencia",
                               "estado_verificacion"], start=1):
            v = x.get(k)
            wn.cell(i, j, str(v) if isinstance(v, dict) else v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
